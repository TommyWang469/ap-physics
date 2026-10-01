"""Builds everything in turn-in/ from the Logger Pro .CMBL files in this folder.

Run:  python3 make_report.py
"""
import math
import re
import shutil
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).parent
OUT = HERE / "turn-in"
G = 9.8
M_BLOCK = 0.462  # kg, measured on the balance (data sheet, Part I)
BEST2, BEST3 = "part2trial1", "part3trial1"


def sig4(x):
    """Round the way Logger Pro displays it (4 significant figures)."""
    return float(f"{x:.4g}")


def load(name):
    s = (HERE / f"{name}.CMBL").read_text(encoding="utf-8", errors="replace")
    cols = [np.array(c.split(), float)
            for c in re.findall(r"<ColumnCells>\s*(.*?)</ColumnCells>", s, re.S)]
    return s, cols


def part2(name):
    """Peak force and mean force over the Statistics ranges saved in the file."""
    s, (t, F) = load(name)
    lo = [float(x) for x in re.findall(r"<GUIHelperRangeMin>(.*?)<", s)]
    hi = [float(x) for x in re.findall(r"<GUIHelperRangeMax>(.*?)<", s)]
    # peak is taken inside the first range (the breakaway), not the whole run:
    # part2.1trial1 has a later bump while sliding that is higher than the breakaway
    start = (t >= lo[0]) & (t <= hi[0])
    kin = (t >= lo[1]) & (t <= hi[1])
    i = np.flatnonzero(start)[F[start].argmax()]
    return dict(t=t, F=F, peak=sig4(F[i]), t_peak=t[i],
                kin=sig4(F[kin].mean()), k0=lo[1], k1=hi[1])


def part3(name):
    """Slope of the Linear Fit shown on the velocity graph, as saved in the file."""
    s, (t, x) = load(name)
    h = re.search(r"<PageHelperDataCurveFit>(.*?)</PageHelperDataCurveFit>", s, re.S).group(1)
    get = lambda k: re.search(f"<{k}>(.*?)</{k}>", h).group(1).strip()
    fid = get("CurveFitFunctionID")
    m, b = map(float, re.search(
        rf"<ID>{fid} </ID>.*?<FunctionCoefficientArray>2 (\S+) (\S+)", s, re.S).groups())
    return dict(t=t, x=x, a=sig4(m), b=b, corr=float(get("CurveFitCorrelation")),
                f0=float(get("GUIHelperRangeMin")), f1=float(get("GUIHelperRangeMax")))


def linfit(x, y):
    m, b = np.polyfit(x, y, 1)
    return m, b


# ---------------------------------------------------------------- numbers
added = [0.5, 1.0, 1.5]                       # kg on top of the block in Part II
files2 = ["part2", "part2.1", "part2.2"]
mass2 = [round(M_BLOCK + a, 3) for a in added]
N2 = [m * G for m in mass2]
runs2 = [[part2(f"{f}trial{i}") for i in (1, 2, 3)] for f in files2]
peak = [[r["peak"] for r in row] for row in runs2]
kin = [[r["kin"] for r in row] for row in runs2]
peak_avg = [sum(r) / 3 for r in peak]
kin_avg = [sum(r) / 3 for r in kin]
mu_s, b_s = linfit(N2, peak_avg)
mu_k2, b_k2 = linfit(N2, kin_avg)
ratio_k = [f / n for f, n in zip(kin_avg, N2)]

mass3 = [M_BLOCK, round(M_BLOCK + 0.5, 3)]
runs3 = [[part3(f"{f}trial{i}") for i in (1, 2, 3)] for f in ("part3", "part3.1")]
acc = [[r["a"] for r in row] for row in runs3]
fk3 = [[m * a for a in row] for m, row in zip(mass3, acc)]
mu3 = [[a / G for a in row] for row in acc]
mu3_avg = [sum(r) / 3 for r in mu3]
fk3_avg = [sum(r) / 3 for r in fk3]

b2 = runs2[0][0]          # best Part II run
b3 = runs3[0][0]          # best Part III run
v0, v1 = (abs(b3["a"] * t + b3["b"]) for t in (b3["f0"], b3["f1"]))
theta = math.degrees(math.atan(mu_s))
mu_lo, mu_hi = (math.tan(math.radians(theta + d)) for d in (-2, 2))


# ---------------------------------------------------------------- graphs
def fit_graph(y, m, b, ylabel, title, sym, fname):
    fig, ax = plt.subplots(figsize=(6.5, 4.5))
    xs = np.array([0, 21])
    ax.plot(xs, m * xs + b, "k-", lw=1)
    ax.plot(N2, y, "o", color="#c0392b", ms=8)
    ax.set(xlim=(0, 21), ylim=(min(-3, b - 1), 13), xlabel="Normal force (N)",
           ylabel=ylabel, title=title)
    ax.axhline(0, color="gray", lw=0.6)
    ax.grid(alpha=0.3)
    sign = "+" if b >= 0 else "-"
    ax.text(0.04, 0.93, f"y = {m:.3f}x {sign} {abs(b):.2f}\nslope = {sym} = {m:.3f}",
            transform=ax.transAxes, va="top", bbox=dict(fc="white", ec="gray"))
    fig.tight_layout()
    fig.savefig(OUT / fname, dpi=200)
    plt.close(fig)


def velocity_graph(fname):
    t, x = b3["t"], b3["x"]
    v = np.gradient(x, t)            # Logger Pro does not save its velocity column
    fig, ax = plt.subplots(figsize=(6.5, 4))
    ax.plot(t, v, color="#c0392b")
    ax.axvspan(b3["f0"], b3["f1"], color="gold", alpha=0.35)
    tt = np.array([b3["f0"], b3["f1"]])
    ax.plot(tt, b3["a"] * tt + b3["b"], "k-", lw=1.5)
    ax.set(xlim=(0.3, 1.4), ylim=(-1.15, 0.5), xlabel="Time (s)", ylabel="Velocity (m/s)",
           title="Part III, no added mass, trial 1")
    ax.annotate("at rest", (0.45, 0), (0.36, 0.25), arrowprops=dict(arrowstyle="->"))
    ax.annotate("push", (0.65, -0.6), (0.36, -0.75), arrowprops=dict(arrowstyle="->"))
    ax.annotate(f"sliding freely (shaded)\nslope = a = {b3['a']:.3f} m/s²",
                (0.85, b3["a"] * 0.85 + b3["b"]), (0.95, -0.9),
                arrowprops=dict(arrowstyle="->"))
    ax.annotate("stopped", (1.1, 0), (1.15, 0.25), arrowprops=dict(arrowstyle="->"))
    ax.axhline(0, color="gray", lw=0.6)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(OUT / fname, dpi=200)
    plt.close(fig)


def fbd(fname):
    fig, ax = plt.subplots(figsize=(5, 3.6))
    ax.add_patch(plt.Rectangle((-0.6, -0.35), 1.2, 0.7, fc="#e8d5b0", ec="k"))
    ax.plot([-2.6, 2.6], [-0.35, -0.35], "k-", lw=2)
    arrow = dict(arrowstyle="-|>", lw=2, color="k")
    for (dx, dy), label, (lx, ly) in [((0, 1.5), "N (normal force)", (0.1, 1.55)),
                                      ((0, -1.5), "mg (weight)", (0.1, -1.6)),
                                      ((1.6, 0), "f_k (kinetic friction)", (0.75, 0.15))]:
        ax.annotate("", (dx, dy), (0, 0), arrowprops=arrow)
        ax.text(lx, ly, label, fontsize=11)
    ax.annotate("", (-2.3, 0.75), (-1.0, 0.75),
                arrowprops=dict(arrowstyle="-|>", ls="--", color="gray"))
    ax.text(-2.5, 0.9, "motion (toward detector)", color="gray", fontsize=9)
    ax.set(xlim=(-2.7, 2.9), ylim=(-1.9, 1.9))
    ax.axis("off")
    fig.tight_layout()
    fig.savefig(OUT / fname, dpi=200)
    plt.close(fig)


# ---------------------------------------------------------------- labeled Logger Pro printouts
def label_image(src, dst, notes):
    """notes: (text, text_xy, arrow_tip_xy) in fractions of the image size."""
    im = Image.open(HERE / src).convert("RGB")
    W, H = im.size
    d = ImageDraw.Draw(im)
    font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", int(H * 0.024))
    blue = (20, 60, 200)
    for text, (tx, ty), tip in notes:
        x, y = tx * W, ty * H
        box = d.multiline_textbbox((x, y), text, font=font)
        pad = H * 0.006
        d.rectangle((box[0] - pad, box[1] - pad, box[2] + pad, box[3] + pad),
                    fill="white", outline=blue, width=3)
        d.multiline_text((x, y), text, font=font, fill=blue)
        if tip:
            px, py = tip[0] * W, tip[1] * H
            sx = min(max(px, box[0]), box[2])
            sy = box[1] - pad if py < box[1] else box[3] + pad
            d.line((sx, sy, px, py), fill=blue, width=5)
            r = H * 0.007
            d.ellipse((px - r, py - r, px + r, py + r), fill=blue)
    im.save(OUT / dst)


def labeled_printouts():
    label_image(f"{BEST2}.png", "PRINT part2 force vs time (labeled).png", [
        ("Block at rest\n(not pulling yet)", (0.10, 0.74), (0.20, 0.925)),
        ("Still at rest: static friction\ngrows to match my pull", (0.42, 0.74), (0.355, 0.70)),
        (f"Block just starts to move\npeak static friction = {b2['peak']:.3f} N", (0.50, 0.40), (0.394, 0.498)),
        (f"Moving at constant speed\nkinetic friction = {b2['kin']:.3f} N (mean)", (0.60, 0.64), (0.70, 0.565)),
    ])
    label_image(f"{BEST3}.png", "PRINT part3 velocity vs time (labeled).png", [
        ("Block sliding freely toward the detector and slowing down.\n"
         "Friction is the only horizontal force here.", (0.30, 0.86), (0.25, 0.775)),
        (f"Straight line = constant acceleration\nslope = a = {b3['a']:.3f} m/s²", (0.62, 0.58), (0.70, 0.705)),
    ])


# ---------------------------------------------------------------- report
def build_docx():
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)

    def para(text="", bold_lead=None):
        p = doc.add_paragraph()
        if bold_lead:
            p.add_run(bold_lead + " ").bold = True
        p.add_run(text)
        return p

    def pic(name, width=5.8):
        doc.add_picture(str(OUT / name), width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(header, rows):
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for c, h in zip(t.rows[0].cells, header):
            c.text = h
            c.paragraphs[0].runs[0].bold = True
        for r in rows:
            for c, v in zip(t.add_row().cells, r):
                c.text = v
        doc.add_paragraph()

    f3 = lambda x: f"{x:.3f}"
    f4 = lambda x: f"{x:#.4g}"

    doc.add_heading("Static and Kinetic Friction Lab", 0)
    para("Tommy Wang")

    doc.add_heading("Preliminary Questions", 1)
    para("The force to start the box moving is greater than the force to keep it moving. "
         "I'm basing this on pushing heavy furniture. It takes a hard shove to get it going, "
         "and then it is easier to keep it sliding.", "1.")
    para("I think friction gets bigger when the box is heavier, and that it goes up in "
         "proportion to the weight. A heavier box pushes down on the floor harder, so the "
         "surfaces are pressed together more.", "2.")

    doc.add_heading("Data Tables", 1)
    para(f"Part I: mass of block = {M_BLOCK} kg")
    para("Part II: Peak static friction")
    table(["Total mass (kg)", "Normal force (N)", "Trial 1 (N)", "Trial 2 (N)", "Trial 3 (N)", "Average (N)"],
          [[f3(m), f4(n), *map(f4, p), f4(a)] for m, n, p, a in zip(mass2, N2, peak, peak_avg)])
    para("Part II: Kinetic friction")
    table(["Total mass (kg)", "Normal force (N)", "Trial 1 (N)", "Trial 2 (N)", "Trial 3 (N)", "Average (N)"],
          [[f3(m), f4(n), *map(f4, k), f4(a)] for m, n, k, a in zip(mass2, N2, kin, kin_avg)])
    for label, m, arow, frow, murow, avg in zip(
            ["Part III: Block with no additional mass", "Part III: Block with 500 g additional mass"],
            mass3, acc, fk3, mu3, mu3_avg):
        para(f"{label} (m = {m} kg)")
        table(["Trial", "Acceleration (m/s²)", "Kinetic friction force (N)", "μk"],
              [[str(i + 1), f4(a), f3(f), f3(mu)] for i, (a, f, mu) in enumerate(zip(arow, frow, murow))]
              + [["", "", "Average μk:", f3(avg)]])

    doc.add_heading("Analysis", 1)
    para("Labeled force vs. time graph (500 g added, trial 1):", "1.")
    pic("PRINT part2 force vs time (labeled).png", 6.3)
    para(f"The block is at rest from 0 s until the peak at {b2['t_peak']:.2f} s. I start pulling at "
         f"about 1.2 s, and the force climbs while the block still doesn't move. The block just "
         f"starts to move at the peak, {b2['peak']:.3f} N at {b2['t_peak']:.2f} s. After that it moves at "
         f"constant speed from about {b2['k0']:.2f} s to {b2['k1']:.2f} s, where the force is flat "
         f"at {b2['kin']:.3f} N.")

    para(f"It took {b2['peak']:.3f} N to start the block sliding but only {b2['kin']:.3f} N to keep it "
         f"sliding, so starting takes more force. This matches my experience. A heavy box on the "
         f"floor is hardest to push right at the start, and once it breaks loose it is easier to "
         f"keep going.", "2.")

    para(f"Greater than. With the same normal force, the peak static friction ({b2['peak']:.3f} N) was "
         f"bigger than the kinetic friction ({b2['kin']:.3f} N). Since the coefficient is friction "
         f"divided by normal force, the static coefficient has to be bigger too.", "3.")

    para("The table is horizontal, so the normal force equals the weight, N = mg. "
         + " ".join(f"{m:.3f} kg × 9.8 = {f4(n)} N." for m, n in zip(mass2, N2))
         + " These are filled in for both Part II tables.", "4.")

    para("Maximum static friction vs. normal force:", "5.")
    pic("graph 5 static friction vs normal force.png")

    para(f"The slope is {mu_s:.3f}, so μs = {mu_s:.3f}. It has no units, because it is newtons "
         f"divided by newtons. In theory the line should pass through the origin, because "
         f"F = μs·N means zero normal force gives zero friction. My best-fit line is "
         f"y = {mu_s:.3f}x − {abs(b_s):.2f}, so it misses the origin by {abs(b_s):.2f} N. I only have three "
         f"points and they don't sit on one straight line. The peak depends on how smoothly I "
         f"pulled, and at {mass2[1]:.3f} kg my three peaks went from {min(peak[1]):.3f} N to "
         f"{max(peak[1]):.3f} N.", "6.")

    para("Average kinetic friction vs. normal force:", "7.")
    pic("graph 7 kinetic friction vs normal force.png")
    para(f"The slope is {mu_k2:.3f}, so μk = {mu_k2:.3f} from this graph. It should pass through the "
         f"origin for the same reason as before: F = μk·N, so no normal force means no friction. "
         f"My line is y = {mu_k2:.3f}x + {b_k2:.2f}, so it does not. It crosses the y-axis at {b_k2:.2f} N. "
         f"The kinetic friction went up with normal force, but not in proportion. Dividing each "
         f"average by its normal force gives {ratio_k[0]:.3f}, {ratio_k[1]:.3f} and {ratio_k[2]:.3f}, "
         f"so the heaviest run came out low and that flattens the slope.")

    para("Free-body diagram for the sliding block:", "8.")
    pic("free body diagram.png", 3.8)
    para(f"After the block leaves my hand, friction is the only horizontal force, so f_k = ma. "
         f"Vertically, N = mg. Example for trial 1 with no added mass: "
         f"f_k = {M_BLOCK} × {acc[0][0]:.3f} = {fk3[0][0]:.3f} N. The other trials are in the data table.")
    para("Sample velocity vs. time graph (no added mass, trial 1). The shaded part is the section "
         "I used for the acceleration:")
    pic("part3 velocity whole run.png")
    pic("PRINT part3 velocity vs time (labeled).png", 6.3)
    para(f"Logger Pro's linear fit on that section gives a slope of {b3['a']:.3f} m/s², "
         f"with correlation {b3['corr']:.4f}.")

    para(f"μk = f_k / N = ma / mg = a / g. Example: {acc[0][0]:.3f} / 9.8 = {mu3[0][0]:.3f}. "
         f"The average is {mu3_avg[0]:.3f} for the block alone and {mu3_avg[1]:.3f} for the block "
         f"with 500 g. All six values are in the data table.", "9.")

    para(f"No. In the trial 1 graph the block slowed from about {v0:.1f} m/s to about {v1:.1f} m/s, "
         f"and the velocity graph stayed a straight line the whole way (correlation {b3['corr']:.4f}). "
         f"A straight line means the acceleration was constant, so the friction force was the same "
         f"at high speed and at low speed. If μk depended on speed the line would curve.", "10.")

    para(f"Yes. In Part II the average kinetic friction went from {f4(kin_avg[0])} N to "
         f"{f4(kin_avg[1])} N to {f4(kin_avg[2])} N as the normal force went from {f4(N2[0])} N to "
         f"{f4(N2[1])} N to {f4(N2[2])} N. In Part III the average friction force was "
         f"{fk3_avg[0]:.3f} N for the block alone and {fk3_avg[1]:.3f} N with 500 g added. "
         f"More weight gave more friction every time.", "11.")

    pct = (mu3_avg[1] - mu3_avg[0]) / mu3_avg[0] * 100
    para(f"Not really. In Part III the weight more than doubled ({M_BLOCK} kg to {mass3[1]} kg) but "
         f"μk only went from {mu3_avg[0]:.3f} to {mu3_avg[1]:.3f}, about {pct:.0f}% different. In Part II, "
         f"friction divided by normal force was {ratio_k[0]:.3f}, {ratio_k[1]:.3f} and {ratio_k[2]:.3f}. "
         f"The first two agree with Part III. The heaviest one is lower, but I think that is from "
         f"pulling by hand and not a real change, since Part III doesn't show μk going down with "
         f"weight.", "12.")

    para(f"Part III gave μk = {mu3_avg[0]:.3f} (block alone) and {mu3_avg[1]:.3f} (with 500 g). "
         f"Part II gave μk = {mu_k2:.3f} from the slope. I expected them to be the same, because it "
         f"is the same block on the same surface and μk only depends on the two surfaces. They are "
         f"not the same. The Part II slope is small because the best-fit line has a {b_k2:.2f} N "
         f"intercept and doesn't go through the origin. If I use friction divided by normal force "
         f"for each Part II mass, I get {ratio_k[0]:.3f} and {ratio_k[1]:.3f} for the two lighter "
         f"loads, which do agree with Part III. I trust Part III more, because the acceleration "
         f"comes from the motion detector and a linear fit, while Part II depends on me pulling at "
         f"a steady speed by hand.", "13.")

    doc.add_heading("Extension", 1)
    para(f"On an incline the block starts to slide when mg·sinθ = μs·mg·cosθ, so tanθ = μs. "
         f"With μs = {mu_s:.3f}, θ = arctan({mu_s:.3f}) = {theta:.1f}°. If the angle was measured "
         f"2° off, I would get μs = tan({theta - 2:.1f}°) = {mu_lo:.3f} or tan({theta + 2:.1f}°) = {mu_hi:.3f}. "
         f"So the coefficient would change by about {(mu_hi - mu_lo) / 2:.2f}, which is about "
         f"{(mu_hi - mu_lo) / 2 / mu_s * 100:.0f}%.", "1.")

    doc.save(OUT / "Friction Lab Report.docx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    up = OUT / "UPLOAD online"
    up.mkdir(exist_ok=True)
    for name in (BEST2, BEST3):
        shutil.copy2(HERE / f"{name}.CMBL", up)
        shutil.copy2(HERE / f"{name}.png", up)
    fit_graph(peak_avg, mu_s, b_s, "Average peak static friction (N)",
              "Peak Static Friction vs. Normal Force", "μs", "graph 5 static friction vs normal force.png")
    fit_graph(kin_avg, mu_k2, b_k2, "Average kinetic friction (N)",
              "Kinetic Friction vs. Normal Force", "μk", "graph 7 kinetic friction vs normal force.png")
    velocity_graph("part3 velocity whole run.png")
    fbd("free body diagram.png")
    labeled_printouts()
    build_docx()

    # self-check: the same quantity must agree everywhere it is used
    assert abs(b3["a"] - 4.136) < 5e-4 and abs(b2["peak"] - 4.847) < 5e-4
    assert all(abs(f - m * a) < 1e-9 for m, row, fr in zip(mass3, acc, fk3) for a, f in zip(row, fr))
    print("peak", peak, [f"{x:.4f}" for x in peak_avg])
    print("kin ", kin, [f"{x:.4f}" for x in kin_avg])
    print(f"mu_s {mu_s:.4f} b {b_s:.3f} | mu_k2 {mu_k2:.4f} b {b_k2:.3f} | ratios", [f"{r:.3f}" for r in ratio_k])
    print("acc", acc, "fk", [[f"{f:.3f}" for f in r] for r in fk3])
    print("mu3", [[f"{m:.3f}" for m in r] for r in mu3], [f"{m:.4f}" for m in mu3_avg], [f"{f:.3f}" for f in fk3_avg])
    print(f"v {v0:.2f}->{v1:.2f} theta {theta:.2f} mu {mu_lo:.3f} {mu_hi:.3f} kin range {b2['k0']:.2f}-{b2['k1']:.2f}")
